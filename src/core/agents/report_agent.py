"""Bounded section reports with explicit grounding and coverage."""

import asyncio
import inspect
import json
from io import BytesIO

from docx import Document

from core.models.agents import AgentResult
from core.reports.grounding import UNAVAILABLE, VERIFICATION_SCOPE, validate_section
from core.reports.sources import ReportClarificationRequired, authorize_sources, collect_sources, serialize_context
from core.reports.templates import default_report_template, parse_report_template, render_report_template
from core.prompts.report import REPORT_SECTION_PROMPT, REPORT_REPAIR_PROMPT

MIME_TYPES = {
    "markdown": "text/markdown; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ReportAgent:
    def __init__(self, knowledge_agent, analytics_agent, *, llm=None, source_router=None):
        self.knowledge_agent = knowledge_agent
        self.analytics_agent = analytics_agent
        self.llm = llm
        self._owns_llm = llm is None
        self.source_router = source_router

    async def close(self):
        if self._owns_llm and self.llm is not None:
            self.llm.close()

    async def run(self, request):
        result, payload, _ = await self.generate_from_template(
            request, source_mode=request.metadata.get("source_mode", "documents"))
        if payload:
            result.answer = payload.decode("utf-8")
        return result

    async def generate_from_template(self, request, filename=None, content=None, *,
                                     source_mode="documents", output_format=None):
        authorize_sources(request, source_mode)
        if output_format is not None and output_format not in MIME_TYPES:
            raise ValueError("Output format must be markdown or docx")
        template = (parse_report_template(filename, content) if filename is not None
                    else default_report_template(output_format or "markdown"))
        if output_format and template.format != output_format:
            raise ValueError("Output format conflicts with uploaded template")
        if source_mode in {'orchestrated', 'combined'} and self.source_router is None:
            from core.agents.analytics_agent import AnalyticsAgent
            from core.agents.router import AgentRouter
            analytics = (self.analytics_agent if isinstance(self.analytics_agent, AnalyticsAgent)
                         else AnalyticsAgent(self.analytics_agent))
            self.source_router = AgentRouter(None, self.knowledge_agent, analytics, llm=self._provider())
        try:
            contexts, coverage, limitations, row_coverage = await collect_sources(
                request, template.headings, source_mode, self.knowledge_agent, self.analytics_agent,
                source_router=self.source_router)
        except ReportClarificationRequired as error:
            clarification = str(error)
            return AgentResult(answer=clarification, agent="report", metadata={
                "ok": False, "status": "clarification_required", "clarification_required": True,
                "limitations": [clarification],
            }), b"", MIME_TYPES[template.format]
        sections = {}
        tables = [[] for _ in template.table_headers]
        section_status = {}
        issues = {}
        evidence = {}
        usable = 0
        provider_failed = False
        for heading in template.headings:
            supplied = contexts[heading]
            indices = [index for index, owner in enumerate(template.table_sections) if owner == heading]
            headers = [template.table_headers[index] for index in indices]
            body, rows = UNAVAILABLE, [[] for _ in headers]
            valid = False
            if supplied:
                prompt = self._prompt(request.question, heading, headers, supplied)
                for _ in range(2):
                    try:
                        raw = await asyncio.wait_for(self._generate(prompt), timeout=90)
                    except Exception:
                        provider_failed = True
                        issues[heading] = "provider_unavailable"
                        break
                    try:
                        draft = self._decode_draft(raw)
                        body, rows = validate_section(draft, heading, headers, supplied)
                        valid = body != UNAVAILABLE or any(rows)
                        if valid:
                            break
                    except (ValueError, TypeError, KeyError, RecursionError) as error:
                        issues[heading] = str(error)[:100]
                        prompt += REPORT_REPAIR_PROMPT.format(issue=issues[heading])
            if valid:
                usable += 1
                issues.pop(heading, None)
                section_status[heading] = "complete"
                for item in supplied:
                    evidence[item.citation_id] = item
            else:
                body, rows = UNAVAILABLE, [[] for _ in headers]
                section_status[heading] = "failed"
                limitations.append(f"Section unavailable or verification failed: {heading}")
            sections[heading] = body
            for index, values in zip(indices, rows):
                tables[index] = values
        limitations = list(dict.fromkeys(limitations))
        status = "failed" if not usable else ("partial" if limitations or usable < len(sections) else "complete")
        metadata = {"ok": status == "complete", "status": status, "limitations": limitations,
                    "sections": section_status, "coverage": coverage, "row_coverage": row_coverage,
                    "verification_scope": VERIFICATION_SCOPE, "issues": issues,
                    "failure_type": ("source_unavailable" if not any(contexts.values()) else (
                        "provider_unavailable" if provider_failed else "validation_failed")) if status == "failed" else None}
        if 'orchestration' in row_coverage:
            metadata['source_collection'] = row_coverage['orchestration']
        result = AgentResult(answer=f"Report status: {status}.", agent="report",
                             evidence=list(evidence.values()), metadata=metadata,
                             data={"report_format": template.format, "section_count": len(sections)})
        if status == "failed":
            return result, b"", MIME_TYPES[template.format]
        output = render_report_template(template, sections, tables)
        appendix = self._appendix(status, limitations, evidence.values(), coverage, row_coverage)
        if isinstance(output, str):
            payload = (output + appendix).encode("utf-8")
        else:
            document = Document(BytesIO(output))
            document.add_heading("Report status and sources", 1)
            for line in appendix.splitlines():
                if line.strip():
                    document.add_paragraph(line.lstrip("# "))
            stream = BytesIO()
            document.save(stream)
            payload = stream.getvalue()
        return result, payload, MIME_TYPES[template.format]

    async def _generate(self, prompt):
        provider = self._provider()
        options = {"max_output_tokens": 4096} if "max_output_tokens" in inspect.signature(provider.generate).parameters else {}
        raw = await asyncio.to_thread(provider.generate, prompt, **options)
        return await raw if inspect.isawaitable(raw) else raw

    def _provider(self):
        if self.llm is None:
            from core.llm.groq_provider import GroqProvider
            self.llm = GroqProvider()
        return self.llm

    @staticmethod
    def _decode_draft(raw):
        if not isinstance(raw, str) or len(raw) > 100_000:
            raise ValueError("invalid_response_size")
        text = raw.strip()
        if text.startswith("```json") and text.endswith("```"):
            text = text[7:-3].strip()
        return json.loads(text)

    @staticmethod
    def _prompt(question, heading, headers, evidence):
        return REPORT_SECTION_PROMPT.format(
            unavailable=UNAVAILABLE,
            question=json.dumps(question, ensure_ascii=False),
            heading=json.dumps(heading, ensure_ascii=False),
            headers=json.dumps(headers, ensure_ascii=False),
            evidence=serialize_context(evidence),
        )

    @staticmethod
    def _appendix(status, limitations, evidence, coverage, rows):
        lines = ["", "## Report status and sources", f"Status: {status}", VERIFICATION_SCOPE,
                 rows["scope"], f"SQL rows considered: {rows['considered_rows']} of {rows['available_rows']} returned."]
        lines.extend(f"- {note}" for note in limitations)
        for heading, counts in coverage.items():
            lines.append(f"Coverage — {heading}: {counts['included']} included; {counts['omitted']} omitted.")
        for item in evidence:
            page = item.metadata.get("page_number", item.metadata.get("page"))
            location = f"; page {page}" if page is not None else ""
            lines.append(f"[{item.citation_id}] {item.source}{location}")
        return "\n".join(lines) + "\n"
