"""Parse and populate bounded Markdown and DOCX report templates."""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
MAX_DOCX_EXPANDED_BYTES = 20 * 1024 * 1024
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_TABLE_SEPARATOR = re.compile(r"^\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?$")


@dataclass(frozen=True)
class ReportTemplate:
    format: str
    content: str | bytes
    headings: tuple[str, ...]
    table_headers: tuple[tuple[str, ...], ...]
    table_sections: tuple[str, ...] = ()


def default_report_template(output_format="markdown") -> ReportTemplate:
    if output_format not in {"markdown", "docx"}:
        raise ValueError("Output format must be markdown or docx")
    if output_format == "markdown":
        return parse_report_template("report.md", b"# Summary\n\n## Findings\n\n## Limitations\n")
    document = Document()
    for heading in ("Summary", "Findings", "Limitations"):
        document.add_heading(heading, 1)
    raw = BytesIO()
    document.save(raw)
    return parse_report_template("report.docx", raw.getvalue())


def _bounded_template(format, content, headings, tables, owners):
    if not headings or len(headings) > 20:
        raise ValueError("Template must contain between 1 and 20 headings")
    if len(headings) != len(set(headings)):
        raise ValueError("Template headings must be unique")
    if any(len(heading) > 200 for heading in headings) or any(len(cell) > 200 for columns in tables for cell in columns):
        raise ValueError("Heading and table header text must be at most 200 characters")
    if len(tables) > 20 or any(not columns or len(columns) > 20 for columns in tables):
        raise ValueError("Template supports at most 20 tables with 1 to 20 columns")
    if any(not owner for owner in owners):
        raise ValueError("Every table must follow a heading")
    return ReportTemplate(format, content, tuple(headings), tuple(tables), tuple(owners))


def parse_report_template(filename: str, content: bytes) -> ReportTemplate:
    if len(content) > MAX_TEMPLATE_BYTES:
        raise ValueError("Template exceeds the 10 MiB upload limit")
    extension = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if extension == "md":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("Markdown templates must be UTF-8") from error
        headings = tuple(match.group(2).strip() for line in text.splitlines()
                         if (match := _HEADING.match(line)))
        tables = _markdown_table_headers(text)
        owner = ""
        owners = []
        lines = text.splitlines()
        for index, line in enumerate(lines):
            if match := _HEADING.match(line):
                owner = match.group(2).strip()
            if index + 1 < len(lines) and line.strip().startswith("|") and _TABLE_SEPARATOR.match(lines[index + 1].strip()):
                owners.append(owner)
        return _bounded_template("markdown", text, headings, tables, owners)
    if extension != "docx":
        raise ValueError("Template must be a DOCX or Markdown file")
    _validate_docx_package(content)
    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise ValueError("DOCX template could not be read") from error
    headings = tuple(paragraph.text.strip() for paragraph in document.paragraphs
                     if paragraph.style.name.startswith("Heading") and paragraph.text.strip())
    table_headers = tuple(tuple(cell.text.strip() for cell in table.rows[0].cells)
                          for table in document.tables if table.rows)
    owner = ""
    owners = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph) and block.style.name.startswith("Heading") and block.text.strip():
            owner = block.text.strip()
        elif not isinstance(block, Paragraph):
            owners.append(owner)
    return _bounded_template("docx", content, headings, table_headers, owners)


def render_report_template(template: ReportTemplate, sections: dict[str, str],
                           tables: list[list[list[str]]]) -> str | bytes:
    if set(sections) != set(template.headings) or any(
        not isinstance(value, str) or not value.strip() for value in sections.values()
    ):
        raise ValueError("Report must provide non-empty content for every template heading")
    if len(tables) != len(template.table_headers):
        raise ValueError("Report table count must match the template")
    for headers, rows in zip(template.table_headers, tables):
        if any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
            raise ValueError("Report table rows must match the template column count")
        if any(not isinstance(cell, str) for row in rows for cell in row):
            raise ValueError("Report table cells must be text")
    if template.format == "markdown":
        return _render_markdown(template, sections, tables)
    return _render_docx(template, sections, tables)


def _markdown_table_headers(text: str) -> tuple[tuple[str, ...], ...]:
    lines = text.splitlines()
    headers = []
    for index, line in enumerate(lines[:-1]):
        if line.strip().startswith("|") and _TABLE_SEPARATOR.match(lines[index + 1].strip()):
            headers.append(tuple(cell.strip() for cell in line.strip().strip("|").split("|")))
    return tuple(headers)


def _render_markdown(template: ReportTemplate, sections: dict[str, str],
                     tables: list[list[list[str]]]) -> str:
    source_lines = str(template.content).splitlines()
    output = []
    index = 0
    table_index = 0
    while index < len(source_lines):
        line = source_lines[index]
        heading = _HEADING.match(line)
        if heading:
            output.append(line)
            body = sections[heading.group(2).strip()]
            output.extend(("", body.strip(), ""))
            index += 1
            continue
        if (index + 1 < len(source_lines) and line.strip().startswith("|")
                and _TABLE_SEPARATOR.match(source_lines[index + 1].strip())):
            headers = template.table_headers[table_index]
            rows = tables[table_index]
            output.append(line)
            output.append(source_lines[index + 1])
            output.extend("| " + " | ".join(cell.replace("|", "&#124;").replace("\r", "").replace("\n", "<br>")
                                           for cell in row) + " |" for row in rows)
            table_index += 1
            index += 2
            while index < len(source_lines) and source_lines[index].strip().startswith("|"):
                index += 1
            continue
        index += 1
    return "\n".join(output).strip() + "\n"


def _render_docx(template: ReportTemplate, sections: dict[str, str],
                 tables: list[list[list[str]]]) -> bytes:
    document = Document(BytesIO(template.content))
    for paragraph in list(document.paragraphs):
        if not paragraph.style.name.startswith("Heading"):
            paragraph._p.getparent().remove(paragraph._p)
    headings = [paragraph for paragraph in document.paragraphs
                if paragraph.style.name.startswith("Heading") and paragraph.text.strip()]
    for paragraph in reversed(headings):
        previous = paragraph._p
        for text in [part.strip() for part in sections[paragraph.text.strip()].splitlines() if part.strip()]:
            element = OxmlElement("w:p")
            previous.addnext(element)
            inserted = Paragraph(element, paragraph._parent)
            if text.startswith(("- ", "* ")):
                inserted.style = "List Bullet"
                text = text[2:]
            for index, part in enumerate(text.split("**")):
                inserted.add_run(part).bold = bool(index % 2)
            previous = element
    for table, rows in zip(document.tables, tables):
        while len(table.rows) > 1:
            table._tbl.remove(table.rows[-1]._tr)
        for values in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = value
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _validate_docx_package(content: bytes) -> None:
    if not content.startswith(b"PK"):
        raise ValueError("DOCX file content is invalid")
    try:
        with ZipFile(BytesIO(content)) as package:
            expanded_size = 0
            for item in package.infolist():
                path = PurePosixPath(item.filename)
                if path.is_absolute() or ".." in path.parts or "vbaProject.bin" in item.filename:
                    raise ValueError("DOCX package contains unsupported content")
                expanded_size += item.file_size
                if expanded_size > MAX_DOCX_EXPANDED_BYTES:
                    raise ValueError("DOCX expands beyond the 20 MiB processing limit")
            if package.testzip() is not None:
                raise ValueError("DOCX file content is invalid")
    except BadZipFile as error:
        raise ValueError("DOCX file content is invalid") from error
