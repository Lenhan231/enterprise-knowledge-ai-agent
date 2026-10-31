import sys
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from document_processing.docling_adapter.digital_born.docx_semantics import apply_docx_semantics


class DocxSemanticsTest(unittest.TestCase):
    def test_word_section_overrides_mapper_context_on_both_table_and_marker(self):
        document = Document()
        document.add_heading("Specifications", level=1)
        document.add_heading("CTL-4100", level=2)
        document.add_paragraph("Explanation before table.")
        document.add_table(rows=1, cols=2)
        content = BytesIO()
        document.save(content)
        result = {"elements": [
            {"text": text, "metadata": {}}
            for text in ["Specifications", "CTL-4100", "Explanation before table."]
        ] + [{"element_type": "table", "table_ref": 0,
              "metadata": {"section_title": "Old section"}}],
            "tables": [{"metadata": {"section_title": "Old section"}, "context": {}}]}
        apply_docx_semantics(result, content.getvalue())
        for metadata in [result["elements"][-1]["metadata"], result["tables"][0]["metadata"]]:
            self.assertEqual(metadata["section_title"], "CTL-4100")
            self.assertEqual(metadata["section_path"], ["Specifications", "CTL-4100"])

    def test_missing_default_style_preserves_paragraph_text(self):
        document = Document()
        document.add_paragraph("Paragraph without a default style")
        normal = document.styles["Normal"]._element
        normal.getparent().remove(normal)
        content = BytesIO()
        document.save(content)
        self.assertIsNone(Document(BytesIO(content.getvalue())).paragraphs[0].style)
        result = {"elements": [{"text": "Paragraph without a default style", "metadata": {}}]}

        apply_docx_semantics(result, content.getvalue())

        self.assertEqual(result["elements"][0]["text"], "Paragraph without a default style")
        self.assertEqual(result["elements"][0]["metadata"]["section_path"], [])

    def test_defined_heading_style_retains_section_structure(self):
        document = Document()
        document.add_heading("Overview", level=1)
        content = BytesIO()
        document.save(content)
        result = {"elements": [{"text": "Overview", "metadata": {}}]}

        apply_docx_semantics(result, content.getvalue())

        self.assertEqual(result["elements"][0]["metadata"]["heading_level"], 1)
        self.assertEqual(result["elements"][0]["metadata"]["section_path"], ["Overview"])
