import json
import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from core.agents.report_agent import ReportAgent
from core.models.agents import AgentRequest, AgentResult, Evidence, VerifiedIdentity


class ReportAgentTest(unittest.IsolatedAsyncioTestCase):
    async def test_generates_docx_in_template_order_with_benchmark_evidence(self):
        knowledge = MagicMock()
        knowledge.retrieve = AsyncMock(return_value=[Evidence(
            source_type="document", source="policy.docx", content="The policy is a draft.")])
        analytics = MagicMock()
        analytics.run = AsyncMock(return_value=AgentResult(
            answer="Revenue is 100.", agent="analytics",
            data={"rows": [{"revenue": 100}]},
            metadata={"ok": False, "limitations": ["Source data is unavailable."]},
            evidence=[Evidence(citation_id="S1", source_type="database", source="commerce.sales")],
        ))
        llm = MagicMock()
        llm.generate.side_effect = ["{truncated", json.dumps({
            "sections": {"Monthly Report": "Revenue 100 [S1]"}, "tables": [],
        }), json.dumps({"sections": {"Summary": "Draft policy [S2]."},
                       "tables": [{"rows": [["Revenue", "100 [S1]"]]}]})]
        router = MagicMock()
        router.collect_sources = AsyncMock(return_value=AgentResult(
            answer="Sources collected", agent="hybrid",
            data={"tasks": {"s": {"agent": "sql", "status": "partial",
                                   "rows": [{"revenue": 100}], "sources": ["commerce.sales"]}}},
            evidence=[Evidence(source_type="document", source="policy.docx", content="The policy is a draft.",
                               metadata={"task_agent": "knowledge"})],
            metadata={"ok": False, "status": "partial", "limitations": ["Source data is unavailable."]}))
        agent = ReportAgent(knowledge, analytics, llm=llm, source_router=router)
        document = Document()
        document.add_heading("Monthly Report", 1)
        document.add_heading("Summary", 2)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Metric"
        table.rows[0].cells[1].text = "Value"
        source = BytesIO()
        document.save(source)

        request = AgentRequest(question="Create a report", identity=VerifiedIdentity(
            subject="analyst", permissions={"analytics:read", "documents:read"}, document_roles={"employee"}))
        result, payload, media_type = await agent.generate_from_template(
            request, "template.docx", source.getvalue(), source_mode="combined")

        rendered = Document(BytesIO(payload))
        self.assertEqual(media_type, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        self.assertEqual([p.text for p in rendered.paragraphs if p.style.name.startswith("Heading")],
                         ["Monthly Report", "Summary", "Report status and sources"])
        self.assertEqual(rendered.tables[0].rows[1].cells[0].text, "Revenue")
        self.assertEqual([item.source for item in result.evidence], ["Validated analytics rows", "policy.docx"])
        self.assertFalse(result.metadata["ok"])
        self.assertIn("Source data is unavailable.", result.metadata["limitations"])
        self.assertIn("The policy is a draft.", llm.generate.call_args.args[0])
        self.assertEqual(llm.generate.call_count, 3)


if __name__ == "__main__":
    unittest.main()
