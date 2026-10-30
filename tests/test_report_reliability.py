import json
import sys
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock
from io import BytesIO
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from core.agents.report_agent import ReportAgent
from core.models.agents import AgentRequest, AgentResult, Evidence, VerifiedIdentity
from core.reports.sources import collect_sources, serialize_context


def request(analytics=False):
    permissions = {"documents:read"} | ({"analytics:read"} if analytics else set())
    return AgentRequest(question="Revenue report", identity=VerifiedIdentity(
        subject="reader", permissions=permissions, document_roles={"employee"}))


class ReportReliabilityTest(unittest.IsolatedAsyncioTestCase):
    async def test_sql_clarification_stops_report_and_preserves_question(self):
        agent, knowledge, analytics, llm = self.agents()
        clarification = "For returns, do you mean refunded amount, returned units, or return rate?"
        agent.source_router = MagicMock()
        agent.source_router.collect_sources = AsyncMock(return_value=AgentResult(
            answer=clarification, agent="hybrid", metadata={"ok": False, "clarification_required": True}))
        result, output, _ = await agent.generate_from_template(request(True), source_mode="combined")
        self.assertEqual(result.answer, clarification)
        self.assertTrue(result.metadata["clarification_required"])
        self.assertEqual(result.metadata["status"], "clarification_required")
        self.assertEqual(output, b"")
        knowledge.retrieve.assert_not_awaited()
        analytics.run.assert_not_awaited()
        llm.generate.assert_not_called()

    def agents(self, text="Revenue is 100."):
        knowledge = MagicMock()
        knowledge.retrieve = AsyncMock(return_value=[Evidence(
            source_type="document", source="policy.md", content=text)])
        analytics = MagicMock()
        analytics.run = AsyncMock(return_value=AgentResult(
            answer="done", agent="sql", data={"rows": [{"revenue": 100}]},
            metadata={"ok": True}))
        llm = MagicMock()
        llm.generate.return_value = json.dumps({"sections": {"Summary": "Revenue is 100. [S1]"}, "tables": []})
        return ReportAgent(knowledge, analytics, llm=llm), knowledge, analytics, llm

    async def test_document_report_does_not_need_or_query_analytics(self):
        agent, knowledge, analytics, _ = self.agents()
        result, output, _ = await agent.generate_from_template(request(), "a.md", b"# Summary")
        self.assertEqual(result.metadata["status"], "complete")
        self.assertIn(b"Revenue is 100", output)
        analytics.run.assert_not_awaited()
        self.assertTrue(knowledge.retrieve.await_args.args[0].identity.document_roles)

    async def test_denied_analytics_does_no_work(self):
        agent, knowledge, analytics, llm = self.agents()
        with self.assertRaises(PermissionError):
            await agent.generate_from_template(request(), "a.md", b"# Summary", source_mode="combined")
        knowledge.retrieve.assert_not_awaited()
        analytics.run.assert_not_awaited()
        llm.generate.assert_not_called()

    async def test_unknown_citation_and_invented_number_are_never_exported(self):
        agent, _, _, llm = self.agents()
        for text in ("Revenue 999 [S1]", "Revenue 100 [S999]", "Revenue 100"):
            llm.generate.return_value = json.dumps({"sections": {"Summary": text}, "tables": []})
            result, output, _ = await agent.generate_from_template(request(), "a.md", b"# Summary")
            self.assertEqual(result.metadata["status"], "failed")
            self.assertEqual(output, b"")

    async def test_failed_section_is_replaced_without_losing_good_section(self):
        agent, _, _, llm = self.agents()
        llm.generate.side_effect = [
            json.dumps({"sections": {"Summary": "Revenue 100 [S1]"}, "tables": []}),
            "bad JSON", "bad JSON",
        ]
        result, output, _ = await agent.generate_from_template(request(), "a.md", b"# Summary\n# Details")
        self.assertEqual(result.metadata["status"], "partial")
        self.assertIn(b"verification failed", output)
        self.assertIn(b"Status: partial", output)

    async def test_context_not_cut_at_old_2000_character_limit(self):
        agent, _, _, llm = self.agents("Background " * 250 + "Revenue is 100.")
        await agent.generate_from_template(request(), "a.md", b"# Summary")
        self.assertIn("Revenue is 100.", llm.generate.call_args.args[0])

    async def test_sql_rows_are_citable_evidence(self):
        agent, knowledge, _, _ = self.agents()
        result, _, _ = await agent.generate_from_template(
            request(True), "a.md", b"# Summary", source_mode="analytics")
        self.assertEqual(result.metadata["status"], "complete")
        self.assertIn('"revenue": 100', result.evidence[0].content)
        knowledge.retrieve.assert_not_awaited()

    async def test_default_template_and_docx_generation(self):
        agent, _, _, llm = self.agents()
        llm.generate.side_effect = lambda prompt: json.dumps({
            "sections": {json.loads(prompt.split("Heading (data): ")[1].splitlines()[0]): "Revenue 100 [S1]"},
            "tables": []})
        for format in ("markdown", "docx"):
            result, output, _ = await agent.generate_from_template(request(), output_format=format)
            self.assertEqual(result.metadata["status"], "complete")
            if format == "docx":
                document = Document(BytesIO(output))
                self.assertEqual(document.paragraphs[0].text, "Summary")
                self.assertIn("Revenue 100 [S1]", [paragraph.text for paragraph in document.paragraphs])
            else:
                self.assertIn(b"## Findings", output)

    async def test_source_failure_is_partial_and_total_failure_does_not_call_model(self):
        agent, knowledge, _, llm = self.agents()
        knowledge.retrieve.side_effect = RuntimeError("private error")
        result, payload, _ = await agent.generate_from_template(request(), "a.md", b"# Summary")
        self.assertEqual(result.metadata["status"], "failed")
        self.assertEqual(payload, b"")
        llm.generate.assert_not_called()
        self.assertNotIn("private error", str(result.metadata))
        agent.source_router = MagicMock()
        agent.source_router.collect_sources = AsyncMock(return_value=AgentResult(
            answer="Partial sources", agent="hybrid",
            data={"tasks": {"s": {"agent": "sql", "status": "complete", "rows": [{"revenue": 100}]}}},
            metadata={"ok": False, "status": "partial", "limitations": ["Document retrieval unavailable"]}))
        result, payload, _ = await agent.generate_from_template(
            request(True), "a.md", b"# Summary", source_mode="combined")
        self.assertEqual(result.metadata["status"], "partial")
        self.assertIn(b"Document retrieval unavailable", payload)

    async def test_context_budget_and_deduplication_are_disclosed(self):
        agent, knowledge, analytics, _ = self.agents()
        item = Evidence(source_type="document", source="large.md", content="x" * 25000)
        knowledge.retrieve.return_value = [item]
        contexts, coverage, limitations, _ = await collect_sources(
            request(), ["Summary"], "documents", knowledge, analytics)
        self.assertLessEqual(len(serialize_context(contexts["Summary"])), 24000)
        self.assertEqual(coverage["Summary"]["omitted"], 1)
        self.assertTrue(limitations)

    async def test_permission_from_client_metadata_is_ignored(self):
        agent, knowledge, _, _ = self.agents()
        forged = request().model_copy(update={"metadata": {"roles": ["admin"], "permissions": ["analytics:read"]}})
        with self.assertRaises(PermissionError):
            await agent.generate_from_template(forged, source_mode="analytics")
        knowledge.retrieve.assert_not_awaited()

    async def test_generates_more_than_three_grounded_table_rows(self):
        agent, _, _, llm = self.agents("Revenue values 100, 200, 300 and 400.")
        llm.generate.return_value = json.dumps({"sections": {"Summary": "Revenue 100 [S1]"}, "tables": [
            {"rows": [["Revenue", f"{value} [S1]"] for value in (100, 200, 300, 400)]}]})
        result, payload, _ = await agent.generate_from_template(
            request(), "a.md", b"# Summary\n| Metric | Value |\n| --- | --- |")
        self.assertEqual(result.metadata["status"], "complete")
        self.assertIn(b"400 [S1]", payload)

    async def test_wrong_section_citation_is_withheld(self):
        agent, knowledge, _, llm = self.agents()
        knowledge.retrieve.side_effect = [
            [Evidence(source_type="document", source="one.md", content="Revenue 100")],
            [Evidence(source_type="document", source="two.md", content="Revenue 200")]]
        llm.generate.side_effect = [json.dumps({"sections": {"Summary": "Revenue 100 [S1]"}, "tables": []}),
                                   json.dumps({"sections": {"Details": "Revenue 100 [S1]"}, "tables": []}),
                                   json.dumps({"sections": {"Details": "Revenue 100 [S1]"}, "tables": []})]
        result, output, _ = await agent.generate_from_template(request(), "a.md", b"# Summary\n# Details")
        self.assertEqual(result.metadata["status"], "partial")
        self.assertEqual(result.metadata["issues"]["Details"], "unknown_citation")


if __name__ == "__main__":
    unittest.main()
