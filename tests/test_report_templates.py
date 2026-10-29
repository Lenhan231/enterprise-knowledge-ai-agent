import sys
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from core.reports.templates import parse_report_template, render_report_template


class ReportTemplateTest(unittest.TestCase):
    def test_markdown_template_keeps_heading_order_and_fills_table(self):
        source = "# Monthly Report\n\n## Summary\n\n## KPIs\n\n| Metric | Value |\n| --- | --- |\n| placeholder | placeholder |\n"
        template = parse_report_template("report.md", source.encode())
        result = render_report_template(template, {
            "Monthly Report": "December 2022",
            "Summary": "Revenue was supported by the validated rows. [S1]",
            "KPIs": "Results below. [S1]",
        }, [[ ["Revenue", "52,454,936.20"] ]])
        self.assertLess(result.index("# Monthly Report"), result.index("## Summary"))
        self.assertLess(result.index("## Summary"), result.index("## KPIs"))
        self.assertIn("| Revenue | 52,454,936.20 |", result)
        self.assertNotIn("placeholder", result)

    def test_docx_template_preserves_headings_and_table_structure(self):
        source = Document()
        source.add_heading("Monthly Report", 1)
        source.add_heading("Summary", 2)
        table = source.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Metric"
        table.rows[0].cells[1].text = "Value"
        table.rows[1].cells[0].text = "placeholder"
        table.rows[1].cells[1].text = "placeholder"
        raw = BytesIO()
        source.save(raw)
        template = parse_report_template("report.docx", raw.getvalue())
        result = render_report_template(template, {
            "Monthly Report": "December 2022",
            "Summary": "Revenue was supported by the validated rows. [S1]",
        }, [[ ["Revenue", "52,454,936.20"] ]])
        rendered = Document(BytesIO(result))
        self.assertEqual([p.text for p in rendered.paragraphs if p.style.name.startswith("Heading")],
                         ["Monthly Report", "Summary"])
        self.assertEqual(rendered.tables[0].rows[1].cells[0].text, "Revenue")

    def test_rejects_unknown_type_and_oversized_template(self):
        with self.assertRaisesRegex(ValueError, "DOCX or Markdown"):
            parse_report_template("report.pdf", b"content")
        with self.assertRaisesRegex(ValueError, "10 MiB"):
            parse_report_template("report.md", b"x" * (10 * 1024 * 1024 + 1))

    def test_replaces_placeholder_prose_and_escapes_table_cells(self):
        template = parse_report_template("a.md", b"# Summary\nplaceholder prose\n| A | B |\n| --- | --- |\n| old | old |")
        output = render_report_template(template, {"Summary": "New text"}, [[["a|b", "c\nd"]]])
        self.assertNotIn("placeholder", output)
        self.assertNotIn("old", output)
        self.assertIn("a&#124;b", output)
        self.assertIn("c<br>d", output)

    def test_docx_multiline_order_and_placeholder_replacement(self):
        source = Document()
        source.add_heading("Summary", 1)
        source.add_paragraph("placeholder")
        source.add_table(rows=1, cols=1).cell(0, 0).text = "Value"
        source.add_paragraph("another placeholder")
        raw = BytesIO()
        source.save(raw)
        template = parse_report_template("a.docx", raw.getvalue())
        output = render_report_template(template, {"Summary": "First\nSecond\n- Bullet"}, [[["100"]]])
        paragraphs = Document(BytesIO(output)).paragraphs
        self.assertEqual([p.text for p in paragraphs], ["Summary", "First", "Second", "Bullet"])
        self.assertEqual(paragraphs[-1].style.name, "List Bullet")

    def test_rejects_excessive_structure_and_tables_without_owner(self):
        with self.assertRaises(ValueError):
            parse_report_template("a.md", "\n".join(f"# Section {i}" for i in range(21)).encode())
        with self.assertRaisesRegex(ValueError, "follow a heading"):
            parse_report_template("a.md", b"| A | B |\n| --- | --- |\n# Summary")


if __name__ == "__main__":
    unittest.main()
